#!/usr/bin/env python3
"""Generate Word table report from train_loop_log.jsonl (主跑 only, no verify)."""

from __future__ import annotations

import argparse
import json
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

REPO = Path(__file__).resolve().parents[1]
LOG = REPO / "training_platform" / "train_loop_log.jsonl"
RUNS = REPO / "workspace" / "runs"


def run_short(run_id: str | None) -> str:
    if not run_id:
        return "—"
    parts = str(run_id).split("-")
    return parts[1] if len(parts) >= 2 else str(run_id)[:6]


def part_timing_total(run_id: str | None) -> float | None:
    if not run_id:
        return None
    summary = RUNS / run_id / "summary.json"
    if not summary.is_file():
        return None
    data = json.loads(summary.read_text(encoding="utf-8"))
    pt = data.get("part_timing") or {}
    if pt.get("total_s") is not None:
        return float(pt["total_s"])
    return None


def outcome_label(row: dict) -> str:
    report = row.get("report") or {}
    pc = report.get("pixel_comparison") or {}
    dist = pc.get("distance")
    actual = pc.get("actual")
    golden = report.get("golden_pixel") or pc.get("golden")
    if row.get("pixel_within_tolerance"):
        return f"✅ 通过（偏差 {dist}px，actual={actual}，golden={golden}）"
    dx, dy = pc.get("dx"), pc.get("dy")
    return f"❌ 未通过（偏差 {dist}px，dx={dx} dy={dy}，actual={actual}）"


def changes_text() -> str:
    return "无代码改动（train_loop：仅重复跑 agent + auto_analyze，无 Cursor apply / verify）"


def load_rows(
    *,
    session_date: str | None,
    task_contains: str | None,
    run_id_prefix: str | None,
    last_n: int | None,
) -> list[dict]:
    rows: list[dict] = []
    for line in LOG.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        row = json.loads(line)
        st = str(row.get("start_time", ""))
        if session_date and not st.startswith(session_date):
            continue
        rid = str(row.get("run_id", ""))
        if run_id_prefix and not rid.startswith(run_id_prefix):
            continue
        report = row.get("report") or {}
        case_id = report.get("case_id", "")
        if task_contains and task_contains not in case_id and task_contains not in rid:
            # also match via report case_id
            if task_contains.replace("case_", "") not in case_id:
                if "case_002" in (task_contains or "") and case_id != "case_002":
                    continue
                elif "case_011" in (task_contains or "") and case_id != "case_011":
                    continue
        rows.append(row)
    if last_n is not None and last_n > 0:
        rows = rows[-last_n:]
    return rows


def build_doc(
    rows: list[dict],
    *,
    title: str,
    command: str,
    task_desc: str,
    golden_desc: str,
    out_path: Path,
) -> None:
    if not rows:
        raise SystemExit("No matching train_loop rows found.")

    passed = sum(1 for r in rows if r.get("pixel_within_tolerance"))
    times = []
    for r in rows:
        t = part_timing_total(r.get("run_id")) or r.get("agent_elapsed_sec")
        if t:
            times.append(float(t))

    doc = Document()
    h = doc.add_heading(title, level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}\n"
        f"运行命令：{command}\n"
        f"任务：{task_desc}\n"
        f"Golden：{golden_desc}\n"
        "说明：本次为 train_loop（主跑 + auto_analyze；无 verify 阶段）。"
    )

    time_range = ""
    if times:
        time_range = f"；主跑用时 {min(times):.1f}s ~ {max(times):.1f}s"
    doc.add_paragraph(
        f"汇总：共 {len(rows)} 轮；平台像素通过 {passed}/{len(rows)}{time_range}。"
    )

    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    headers = ["轮次", "主跑 run", "主跑用时", "verify run", "verify 用时", "结果", "改动（P0/实际 apply）"]
    for i, hdr in enumerate(headers):
        table.rows[0].cells[i].text = hdr
        for p in table.rows[0].cells[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)

    for row in rows:
        run_id = row.get("run_id")
        psec = part_timing_total(run_id) or row.get("agent_elapsed_sec")
        cells = table.add_row().cells
        cells[0].text = str(row.get("round", "?"))
        cells[1].text = run_short(run_id)
        cells[2].text = f"{float(psec):.1f}s" if psec else "—"
        cells[3].text = "—"
        cells[4].text = "—"
        cells[5].text = outcome_label(row)
        cells[6].text = changes_text()
        for c in cells:
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8)

    doc.add_paragraph("")
    doc.add_paragraph(
        "说明：train_loop 每轮仅主跑 agent 并对比 golden；verify 列填「—」。"
        "若需主跑 + 改码 + verify，请使用 autonomous_loop 并运行 gen_autonomous_rounds_docx.py。"
    )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print(f"Wrote {out_path}")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--date", default="2026-06-30", help="Filter by start_time date prefix")
    parser.add_argument("--case", default="case_002", help="Filter by case_id in report")
    parser.add_argument("--last", type=int, default=None, help="Take last N matching rows")
    parser.add_argument("--run-prefix", default=None, help="Filter run_id prefix e.g. 20260630-175")
    parser.add_argument(
        "--out",
        default="training_platform/case002_train_loop_20260630.docx",
    )
    parser.add_argument(
        "--command",
        default=(
            "python training_platform/train_loop.py --rounds 1 "
            "--task data/cases/case_002_tp332_front/task.yaml"
        ),
    )
    args = parser.parse_args()

    rows = load_rows(
        session_date=args.date,
        task_contains=args.case,
        run_id_prefix=args.run_prefix,
        last_n=args.last,
    )
    # refine case filter
    if args.case:
        rows = [
            r
            for r in rows
            if (r.get("report") or {}).get("case_id", "").startswith(args.case.replace("_tp", ""))
            or args.case in str((r.get("report") or {}).get("case_id", ""))
        ]

    if args.case == "case_002":
        title = "case_002 — train_loop 测试报告（TP332 Step3+ 快速路径）"
        task_desc = "data/cases/case_002_tp332_front/task.yaml（start_from_step3，仅 Part D）"
        golden_desc = "像素 [357, 276]，容差 ±10px（step08_result.json）"
    elif args.case == "case_011":
        title = "case_011 — train_loop 测试报告"
        task_desc = "data/cases/case_011/task.yaml（全量 STANDARD_WORKFLOW）"
        golden_desc = "像素 [2038, 1114] + 6 个 golden 文件，容差 ±10px"
    else:
        title = f"{args.case} — train_loop 测试报告"
        task_desc = args.command
        golden_desc = "见 case_benchmarks.py"

    build_doc(
        rows,
        title=title,
        command=args.command,
        task_desc=task_desc,
        golden_desc=golden_desc,
        out_path=REPO / args.out,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
