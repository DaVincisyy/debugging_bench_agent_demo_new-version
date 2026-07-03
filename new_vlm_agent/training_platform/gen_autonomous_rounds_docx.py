#!/usr/bin/env python3
"""
Generate / update Word table from autonomous_loop rounds (主跑 + verify + 改动).

Usage (after autonomous_loop finishes):
    python training_platform/gen_autonomous_rounds_docx.py
    python training_platform/gen_autonomous_rounds_docx.py --rounds 15
    python training_platform/gen_autonomous_rounds_docx.py --out training_platform/case011_autonomous_15rounds.docx
"""

from __future__ import annotations

import argparse
import json
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

REPO = Path(__file__).resolve().parents[1]
ROUNDS_DIR = REPO / "training_platform" / "autonomous_rounds"
DEFAULT_OUT = REPO / "training_platform" / "case011_autonomous_15rounds.docx"


def run_short(run_id: str | None) -> str:
    if not run_id:
        return "—"
    parts = str(run_id).split("-")
    return parts[1] if len(parts) >= 2 else str(run_id)[:6]


def primary_seconds(row: dict) -> float | None:
    pr = row.get("primary_run") or {}
    pt = pr.get("part_timing") or {}
    if pt.get("total_s") is not None:
        return float(pt["total_s"])
    return None


def verify_seconds(row: dict) -> float | None:
    vr = row.get("verify_run")
    if not vr:
        return None
    pt = vr.get("part_timing") or {}
    if pt.get("total_s") is not None:
        return float(pt["total_s"])
    return None


def outcome_label(row: dict) -> str:
    outcome = str(row.get("apply_outcome") or "")
    sha_after = str(row.get("good_sha_after") or "")[:7]
    pc = (row.get("primary_run") or {}).get("pixel_comparison") or {}
    dist = pc.get("distance_px")

    if outcome == "kept":
        return f"✅ 保留（git {sha_after}）"
    if outcome == "kept_no_commit_diff":
        return f"✅ 保留（无文件 diff，git {sha_after}）"
    if outcome == "reverted_pixel":
        vpx = ((row.get("verify_run") or {}).get("pixel_comparison") or {})
        d = vpx.get("distance_px", dist)
        return f"↩ 回滚（verify 像素未过，偏差 {d}px）"
    if outcome == "reverted_not_session_best":
        return "↩ 回滚（verify 未刷新 session 最快）"
    if outcome == "reverted_no_timing":
        return "↩ 回滚（verify 无耗时数据）"
    if outcome == "skipped_inaccurate_pre_apply":
        return f"↩ 跳过改码（主跑像素未过，偏差 {dist}px）"
    if outcome == "skipped_no_apply_plan":
        return "↩ 跳过（无 apply_plan.json）"
    if outcome == "skipped_empty_actions":
        return "↩ 跳过（apply_plan 无 actions）"
    if outcome == "skipped_guard_blocked_all":
        return "↩ 跳过（guard 拦截全部改动）"
    if outcome == "apply_error":
        return "↩ 异常（Cursor apply 失败）"
    if outcome.startswith("skipped"):
        return f"↩ {outcome}"
    if not (row.get("verify_run") or {}).get("run_id"):
        if dist and not pc.get("within_tolerance"):
            return f"异常/无 verify（主跑像素 {dist}px）"
        return f"异常（{outcome or '无 verify'}）"
    return outcome or "—"


def changes_text(row: dict) -> str:
    items = row.get("improvements_applied") or []
    if not items:
        # reflection-only or skipped
        outcome = row.get("apply_outcome") or ""
        if outcome.startswith("skipped"):
            return "本轮未应用代码改动（见「结果」列原因）。"
        return "无 P0 actions 记录；可能为 reflection-only 或未生成 apply_plan。"
    parts: list[str] = []
    for a in items[:2]:
        title = a.get("title") or a.get("id") or "改动"
        files = a.get("target_files") or []
        fstr = ", ".join(files[:2]) if files else ""
        parts.append(f"P0：{title}" + (f"（{fstr}）" if fstr else ""))
    return "；".join(parts)


def load_rounds(n: int) -> list[dict]:
    rows: list[dict] = []
    for i in range(1, n + 1):
        p = ROUNDS_DIR / f"round_{i:03d}.json"
        if not p.is_file():
            raise FileNotFoundError(f"Missing {p} — run autonomous_loop first.")
        rows.append(json.loads(p.read_text(encoding="utf-8")))
    return rows


def build_doc(rows: list[dict], *, command: str, out_path: Path) -> None:
    doc = Document()
    h = doc.add_heading("case_011 — autonomous_loop 训练报告（主跑 + verify）", level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    kept = sum(1 for r in rows if r.get("apply_outcome") in ("kept", "kept_no_commit_diff"))
    pre_ok = sum(
        1
        for r in rows
        if ((r.get("primary_run") or {}).get("pixel_comparison") or {}).get("within_tolerance")
    )
    doc.add_paragraph(
        f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}\n"
        f"运行命令：{command}\n"
        f"任务：data/cases/case_011/task.yaml | Golden [2038,1114] ±10px\n"
        f"汇总：共 {len(rows)} 轮；主跑像素通过 {pre_ok}/{len(rows)}；代码保留 {kept}/{len(rows)} 轮。"
    )

    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    headers = ["轮次", "主跑 run", "主跑用时", "verify run", "verify 用时", "结果", "改动（P0/实际 apply）"]
    for i, title in enumerate(headers):
        table.rows[0].cells[i].text = title
        for p in table.rows[0].cells[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)

    for r in rows:
        cells = table.add_row().cells
        it = r.get("iteration", "?")
        pr = r.get("primary_run") or {}
        vr = r.get("verify_run") or {}
        psec = primary_seconds(r)
        vsec = verify_seconds(r)

        cells[0].text = str(it)
        cells[1].text = run_short(pr.get("run_id"))
        cells[2].text = f"{psec:.1f}s" if psec is not None else "—"
        cells[3].text = run_short(vr.get("run_id")) if vr else "—"
        cells[4].text = f"{vsec:.1f}s" if vsec is not None else "—"
        cells[5].text = outcome_label(r)
        cells[6].text = changes_text(r)
        for c in cells:
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8)

    doc.add_paragraph("")
    doc.add_paragraph(
        "说明：主跑 = 当前代码跑 agent + 反思；verify = Cursor 改代码后再跑一轮验证；"
        "「保留」表示 verify 通过且刷新 session 最快。"
    )
    doc.save(out_path)
    print(f"Wrote {out_path}")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--rounds", type=int, default=15)
    parser.add_argument("--out", type=str, default=str(DEFAULT_OUT))
    parser.add_argument(
        "--command",
        type=str,
        default="python training_platform/autonomous_loop.py --rounds 15",
    )
    args = parser.parse_args()
    rows = load_rounds(args.rounds)
    build_doc(rows, command=args.command, out_path=Path(args.out))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
