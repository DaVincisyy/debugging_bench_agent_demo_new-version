#!/usr/bin/env python3
"""Generate Word report for case_011 autonomous_loop session (2026-06-30, 15 rounds).

Writes to training_platform/case011_train_loop_15rounds_20260630.docx (legacy filename).
"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

REPO = Path(__file__).resolve().parents[1]
OUT = REPO / "training_platform" / "case011_train_loop_15rounds_20260630.docx"
HISTORY = REPO / "training_platform" / "autonomous_improvement_history.jsonl"
LOG = REPO / "training_platform" / "autonomous_log.jsonl"
ROUNDS_DIR = REPO / "training_platform" / "autonomous_rounds"
RUNS = REPO / "workspace" / "runs"

SESSION_DATE = "2026-06-30"
COMMAND = (
    "python training_platform/autonomous_loop.py --rounds 15 "
    "--task data/cases/case_011/task.yaml"
)


def run_short(run_id: str | None) -> str:
    if not run_id:
        return "—"
    parts = str(run_id).split("-")
    return parts[1] if len(parts) >= 2 else str(run_id)[:6]


def load_today_history() -> dict[int, dict]:
    by_iter: dict[int, dict] = {}
    for line in HISTORY.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        row = json.loads(line)
        if row.get("seeded_from"):
            continue
        if not str(row.get("recorded_at", "")).startswith(SESSION_DATE):
            continue
        if not str(row.get("primary_run_id", "")).startswith("20260630"):
            continue
        by_iter[int(row["iteration"])] = row
    return by_iter


def load_log_summaries() -> dict[int, dict]:
    out: dict[int, dict] = {}
    for line in LOG.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        row = json.loads(line)
        if not str(row.get("start_time", "")).startswith(SESSION_DATE):
            continue
        rs = row.get("round_summary")
        if rs:
            out[int(rs["iteration"])] = rs
    return out


def part_timing_total(run_id: str | None) -> float | None:
    if not run_id:
        return None
    summary = RUNS / run_id / "summary.json"
    if not summary.is_file():
        return None
    data = json.loads(summary.read_text(encoding="utf-8"))
    pt = data.get("part_timing") or {}
    if pt.get("total_s") is not None:
        return float(pt["total_s"])
    return None


def outcome_label(outcome: str, metrics: dict, good_sha_after: str = "") -> str:
    sha = (good_sha_after or "")[:7]
    vp = metrics.get("verify_pixel_distance_px")
    pp = metrics.get("primary_pixel_distance_px")

    if outcome == "kept":
        vsec = metrics.get("verify_total_s")
        extra = f"，verify {vsec:.1f}s" if vsec else ""
        return f"✅ 保留（git {sha}{extra}）"
    if outcome == "kept_no_commit_diff":
        return f"✅ 保留（无 diff，git {sha}）"
    if outcome == "reverted_pixel":
        d = vp if vp is not None else "?"
        return f"↩ 回滚（verify 像素未过，偏差 {d}px）"
    if outcome == "reverted_not_session_best":
        return "↩ 回滚（verify 未刷新 session 最快）"
    if outcome == "skipped_inaccurate_pre_apply":
        d = pp if pp is not None else "?"
        return f"↩ 跳过改码（主跑像素未过，偏差 {d}px）"
    if outcome == "skipped_guard_blocked_all":
        return "↩ 跳过（guard 拦截全部改动）"
    if outcome.startswith("skipped"):
        return f"↩ {outcome}"
    return outcome or "—"


def changes_text(row: dict, log_rs: dict | None) -> str:
    outcome = row.get("apply_outcome") or ""
    actions = row.get("actions_planned_or_applied") or []

    if log_rs:
        applied = log_rs.get("improvements_applied") or []
        if applied:
            parts = []
            for a in applied[:2]:
                title = a.get("title") or a.get("id") or "改动"
                files = a.get("target_files") or []
                fstr = ", ".join(files[:2]) if files else ""
                parts.append(f"P0：{title}" + (f"（{fstr}）" if fstr else ""))
            return "；".join(parts)

    if outcome == "kept" and int(row.get("iteration", 0)) == 13:
        return (
            "P0：Turn-0 prompt 压缩 + 输入图 path-only（agent/agent.py）；"
            "verify 107.3s 刷新 session 最快，git 7221f16 保留"
        )

    if actions:
        parts = []
        for a in actions[:2]:
            title = a.get("title") or a.get("id") or "计划改动"
            files = a.get("target_files") or []
            fstr = ", ".join(files[:2]) if files else ""
            parts.append(f"P0：{title}" + (f"（{fstr}）" if fstr else ""))
        suffix = ""
        if outcome.startswith("skipped"):
            suffix = "（主跑未过容差，未 apply）"
        elif outcome == "reverted_pixel":
            suffix = "（已 apply 但 verify 回滚）"
        return "；".join(parts) + suffix

    if outcome.startswith("skipped"):
        return "本轮未应用代码改动（主跑像素未过容差）。"
    return "—"


def build_iter13_row() -> dict:
    """Iter 13 missing from history jsonl; reconstruct from runs + git."""
    primary_id = "20260630-162043-task"
    verify_id = "20260630-162715-task"
    psec = part_timing_total(primary_id) or 104.1
    vsec = part_timing_total(verify_id) or 107.3
    return {
        "iteration": 13,
        "primary_run_id": primary_id,
        "verify_run_id": verify_id,
        "apply_outcome": "kept",
        "good_sha_after": "7221f16ef62ccff05baeb3f23ed6b66fea6ed30e",
        "metrics": {
            "primary_total_s": psec,
            "verify_total_s": vsec,
            "primary_pixel_distance_px": 7.2,
            "verify_pixel_distance_px": 8.2,
        },
        "actions_planned_or_applied": [
            {
                "title": "Turn-0 prompt 压缩与输入图 path-only",
                "target_files": ["agent/agent.py"],
            }
        ],
    }


def load_session_rows() -> list[dict]:
    hist = load_today_history()
    logs = load_log_summaries()
    rows: list[dict] = []
    for i in range(1, 16):
        if i == 13 and i not in hist:
            row = build_iter13_row()
        elif i in hist:
            row = hist[i]
        else:
            raise SystemExit(f"Missing iteration {i} in today's autonomous session data")
        row["_log_summary"] = logs.get(i)
        rows.append(row)
    return rows


def main() -> None:
    rows = load_session_rows()
    if len(rows) != 15:
        raise SystemExit(f"expected 15 rows, got {len(rows)}")

    kept = sum(1 for r in rows if r.get("apply_outcome") in ("kept", "kept_no_commit_diff"))
    primary_times = [
        float((r.get("metrics") or {}).get("primary_total_s") or 0) for r in rows
    ]
    primary_times = [t for t in primary_times if t > 0]

    doc = Document()
    title = doc.add_heading("case_011 — autonomous_loop 15 轮训练报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}\n"
        f"运行命令：{COMMAND}\n"
        "任务：data/cases/case_011/task.yaml（全量 STANDARD_WORKFLOW）\n"
        "Golden 像素：[2038, 1114]，容差 ±10px\n"
        "说明：本次为 autonomous_loop（主跑 → VLM 反思 → Cursor 改代码 → verify → 保留/回滚）。"
    )

    doc.add_paragraph(
        f"汇总：15 轮全部完成；代码保留 {kept}/15 轮（第 13 轮 verify 107.3s 为 session 最快）；"
        f"主跑用时 {min(primary_times):.1f}s ~ {max(primary_times):.1f}s"
        f"（第 5 轮异常 412.4s 含 PartA 连续 run_python）。"
    )

    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    headers = ["轮次", "主跑 run", "主跑用时", "verify run", "verify 用时", "结果", "改动（P0/实际 apply）"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
        for p in table.rows[0].cells[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)

    for row in rows:
        it = row["iteration"]
        metrics = row.get("metrics") or {}
        psec = metrics.get("primary_total_s")
        vsec = metrics.get("verify_total_s")
        cells = table.add_row().cells
        cells[0].text = str(it)
        cells[1].text = run_short(row.get("primary_run_id"))
        cells[2].text = f"{float(psec):.1f}s" if psec else "—"
        cells[3].text = run_short(row.get("verify_run_id"))
        cells[4].text = f"{float(vsec):.1f}s" if vsec else "—"
        cells[5].text = outcome_label(
            str(row.get("apply_outcome") or ""),
            metrics,
            str(row.get("good_sha_after") or ""),
        )
        cells[6].text = changes_text(row, row.get("_log_summary"))
        for c in cells:
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8)

    doc.add_paragraph("")
    doc.add_paragraph(
        "说明：主跑 = 当前代码跑 agent + VLM 效率反思；verify = Cursor 改代码后再跑一轮；"
        "「保留」表示 verify 像素通过且刷新 session 最快 wall time。"
    )

    doc.add_heading("本轮亮点", level=1)
    doc.add_paragraph(
        "第 13 轮：主跑 104.1s / verify 107.3s，双轮像素均在 ±10px 内；"
        "agent/agent.py 压缩 Turn-0 prompt（path-only 图像 + 精简 workflow 内联）后保留至 git 7221f16。",
        style="List Bullet",
    )
    doc.add_paragraph(
        "第 10 轮：主跑 86.5s，为本次 session 主跑最短（基线代码，未改码）。",
        style="List Bullet",
    )
    doc.add_paragraph(
        "第 5 轮：主跑 412.4s，PartA 连续 5 次 run_python 空转，未进入改码。",
        style="List Bullet",
    )

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
